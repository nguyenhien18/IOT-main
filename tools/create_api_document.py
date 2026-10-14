import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(r"D:\IOT_main\Tai_lieu_API_IoT_Room_Monitoring.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Times New Roman", size=13, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def write_cell(cell, value, bold=False, color=None, size=12):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run(str(value))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def style_table(table, header=True, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_index == 0 and header:
            set_repeat_table_header(row)
        for col_index, cell in enumerate(row.cells):
            if widths and col_index < len(widths):
                cell.width = Cm(widths[col_index])
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=11.5)
            if row_index == 0 and header:
                set_cell_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def add_heading(document, text, level=3):
    p = document.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_paragraph_spacing(p, before=10, after=6, line=1.0)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 3 else 16, bold=True, color=(31, 78, 121))
    return p


def add_label(document, text):
    p = document.add_paragraph()
    p.paragraph_format.keep_with_next = True
    set_paragraph_spacing(p, before=4, after=3, line=1.0)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True)
    return p


def add_note(document, text):
    p = document.add_paragraph()
    set_paragraph_spacing(p, before=2, after=6)
    r1 = p.add_run("Ghi chú: ")
    set_run_font(r1, bold=True, color=(192, 80, 77))
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def add_code(document, value):
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, indent=2)
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F8")
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.0)
    run = p.add_run(value)
    set_run_font(run, name="Consolas", size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_key_value_table(document, rows):
    table = document.add_table(rows=1, cols=2)
    headers = ("Thuộc tính", "Nội dung")
    for idx, value in enumerate(headers):
        write_cell(table.rows[0].cells[idx], value, bold=True, size=11.5)
    for key, value in rows:
        cells = table.add_row().cells
        write_cell(cells[0], key, bold=True, size=11.5)
        write_cell(cells[1], value, size=11.5)
    style_table(table, header=True, widths=[4.2, 11.8])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_parameter_table(document, rows):
    table = document.add_table(rows=1, cols=4)
    for idx, value in enumerate(("Tham số", "Vị trí", "Bắt buộc", "Mô tả")):
        write_cell(table.rows[0].cells[idx], value, bold=True, size=11)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            write_cell(cells[idx], value, bold=(idx == 0), size=10.5)
    style_table(table, header=True, widths=[3.2, 2.5, 2.2, 8.1])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_api(document, index, title, method, endpoint, description, auth=True,
            params=None, request_body=None, response=None, errors=None, note=None):
    add_heading(document, f"3.{index + 1}. API-{index:02d}: {title}")
    add_key_value_table(document, [
        ("API ID", f"API-{index:02d}"),
        ("Method", method),
        ("Endpoint", endpoint),
        ("Base URL", "http://localhost:8080"),
        ("Mô tả", description),
        ("Auth", "Có — Bearer JWT" if auth else "Không cần"),
        ("Content-Type", "application/json"),
    ])

    if auth:
        add_label(document, "Request header:")
        add_code(document, "Authorization: Bearer <jwtToken>\nContent-Type: application/json")

    if params:
        add_label(document, "Tham số:")
        add_parameter_table(document, params)

    if request_body is not None:
        add_label(document, "Request body:")
        add_code(document, request_body)

    add_label(document, "Response 200 OK:")
    add_code(document, response)

    if errors:
        for label, error_body in errors:
            add_label(document, label)
            add_code(document, error_body)

    if note:
        add_note(document, note)


def build_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, before=0, after=6, line=1.0)
    run = title.add_run("TÀI LIỆU MÔ TẢ REST API")
    set_run_font(run, size=18, bold=True, color=(31, 78, 121))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, after=12, line=1.0)
    run = subtitle.add_run("Hệ thống IoT Room Monitoring – Smart Room")
    set_run_font(run, size=14, bold=True)

    add_heading(document, "3.1. Danh sách API")
    overview = [
        ("API-01", "POST", "/api/auth/login", "Đăng nhập và lấy JWT"),
        ("API-02", "GET", "/api/sensors", "Lấy danh sách cảm biến"),
        ("API-03", "GET", "/api/sensors/{id}", "Lấy chi tiết một cảm biến"),
        ("API-04", "GET", "/api/sensors/latest", "Lấy số đo mới nhất"),
        ("API-05", "GET", "/api/data-sensors", "Tìm kiếm, lọc và phân trang dữ liệu cảm biến"),
        ("API-06", "GET", "/api/data-sensors/chart", "Lấy dữ liệu biểu đồ dashboard"),
        ("API-07", "POST", "/api/data-sensors", "Thêm thủ công một số đo cảm biến"),
        ("API-08", "GET", "/api/devices", "Lấy danh sách và trạng thái LED"),
        ("API-09", "POST", "/api/devices/{id}/control", "Điều khiển một LED"),
        ("API-10", "POST", "/api/devices/control-all", "Bật hoặc tắt tất cả LED"),
        ("API-11", "GET", "/api/actions", "Tìm kiếm, lọc và phân trang lịch sử điều khiển"),
        ("API-12", "GET", "/api/actions/{id}", "Lấy chi tiết một lịch sử điều khiển"),
        ("API-13", "GET", "/api/profile", "Lấy thông tin cá nhân"),
        ("API-14", "PUT", "/api/profile", "Cập nhật thông tin cá nhân"),
    ]
    table = document.add_table(rows=1, cols=4)
    for i, value in enumerate(("API ID", "Method", "Endpoint", "Mô tả ngắn")):
        write_cell(table.rows[0].cells[i], value, bold=True, size=11)
    for values in overview:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            write_cell(cells[i], value, bold=(i == 0), size=10.5)
    style_table(table, header=True, widths=[2.2, 2.2, 6.5, 6.0])

    add_note(document, "API-01 và Swagger không cần JWT. Các API còn lại yêu cầu header Authorization: Bearer <jwtToken>. Base URL mặc định là http://localhost:8080.")

    add_api(
        document, 1, "Đăng nhập", "POST", "/api/auth/login",
        "Xác thực tài khoản và trả về JWT dùng cho các API cần đăng nhập.", auth=False,
        request_body={"username": "admin", "password": "123456"},
        response={
            "code": 200, "message": "success",
            "result": {"id": 1, "username": "admin", "role": "ADMIN", "token": "eyJhbGciOiJIUzI1NiJ9...", "type": "Bearer"}
        },
        errors=[("Response 401 Unauthorized:", {"code": 1003, "message": "Invalid username or password"})],
    )

    add_api(
        document, 2, "Lấy danh sách cảm biến", "GET", "/api/sensors",
        "Trả về danh mục các cảm biến được cấu hình trong hệ thống.",
        response={
            "code": 200, "message": "success",
            "result": [
                {"id": 1, "name": "Temperature", "unit": "°C", "createdAt": "2026-10-06T10:00:00"},
                {"id": 2, "name": "Humidity", "unit": "%", "createdAt": "2026-10-06T10:00:00"},
                {"id": 3, "name": "Light", "unit": "lux", "createdAt": "2026-10-06T10:00:00"},
            ]
        },
    )

    add_api(
        document, 3, "Lấy chi tiết một cảm biến", "GET", "/api/sensors/{id}",
        "Lấy thông tin định danh của một cảm biến theo ID; không phải lịch sử số đo.",
        params=[("id", "Path", "Có", "ID cảm biến, ví dụ 1 = Temperature")],
        response={
            "code": 200, "message": "success",
            "result": {"id": 1, "name": "Temperature", "unit": "°C", "createdAt": "2026-10-06T10:00:00"}
        },
        errors=[("Response 404 Not Found:", {"code": 404, "message": "Sensor not found"})],
    )

    add_api(
        document, 4, "Lấy số đo cảm biến mới nhất", "GET", "/api/sensors/latest",
        "Lấy giá trị nhiệt độ, độ ẩm, ánh sáng mới nhất đã được backend nhận và lưu vào MySQL.",
        response={
            "code": 200, "message": "success",
            "result": {"temperature": 28.6, "humidity": 47.0, "light": 5.8, "time": "2026-10-07T15:31:59"}
        },
        note="ESP8266 gửi dữ liệu qua MQTT; frontend gọi API này định kỳ để tạo hiệu ứng dữ liệu thời gian thực.",
    )

    sensor_params = [
        ("sensorId", "Query", "Không", "ID cảm biến: 1, 2 hoặc 3"),
        ("keyword", "Query", "Không", "Tìm theo tên cảm biến, giá trị hoặc thời gian"),
        ("value", "Query", "Không", "Lọc theo giá trị số chính xác"),
        ("time", "Query", "Không", "HH:mm, HH:mm:ss, DD/MM/YYYY hoặc DD/MM/YYYY HH:mm[:ss]"),
        ("from", "Query", "Không", "Thời gian bắt đầu theo ISO-8601"),
        ("to", "Query", "Không", "Thời gian kết thúc theo ISO-8601"),
        ("sort", "Query", "Không", "createdAt,value,id kết hợp asc/desc; mặc định createdAt,desc"),
        ("page", "Query", "Không", "Trang bắt đầu từ 0"),
        ("size", "Query", "Không", "Số bản ghi mỗi trang"),
    ]
    add_api(
        document, 5, "Lấy lịch sử dữ liệu cảm biến", "GET", "/api/data-sensors",
        "Tìm kiếm, lọc, sắp xếp và phân trang lịch sử số đo trực tiếp tại backend.",
        params=sensor_params,
        response={
            "code": 200, "message": "success",
            "result": {
                "content": [{"id": 3318, "sensorId": 1, "sensorName": "Temperature", "value": 28.6, "createdAt": "2026-10-07T15:31:59"}],
                "page": 0, "size": 10, "totalElements": 3318, "totalPages": 332, "first": True, "last": False
            }
        },
        note="Ví dụ: /api/data-sensors?sensorId=1&page=0&size=10&sort=createdAt,desc",
    )

    add_api(
        document, 6, "Lấy dữ liệu biểu đồ dashboard", "GET", "/api/data-sensors/chart",
        "Lấy dữ liệu đã gom theo phút cho ba đường nhiệt độ, độ ẩm và ánh sáng trên dashboard.",
        response={
            "code": 200, "message": "success",
            "result": [
                {"time": "2026-10-07T15:27:00", "temperature": 30.3, "humidity": 72.0, "light": 544.0},
                {"time": "2026-10-07T15:29:00", "temperature": 30.1, "humidity": 71.0, "light": 526.0},
            ]
        },
        note="Backend lấy tối đa 120 bản ghi gần nhất, gom theo phút và trả tối đa 12 mốc thời gian cho biểu đồ.",
    )

    add_api(
        document, 7, "Thêm thủ công một số đo cảm biến", "POST", "/api/data-sensors",
        "Tạo một bản ghi số đo trong MySQL để kiểm thử mà không cần ESP8266.",
        request_body={"sensorId": 1, "value": 28.5, "createdAt": "2026-10-07T15:35:00"},
        response={
            "code": 200, "message": "success",
            "result": {"id": 3319, "sensorId": 1, "sensorName": "Temperature", "value": 28.5, "createdAt": "2026-10-07T15:35:00"}
        },
        note="Có thể bỏ createdAt để backend tự dùng thời gian hiện tại.",
    )

    add_api(
        document, 8, "Lấy danh sách thiết bị", "GET", "/api/devices",
        "Lấy danh sách ba LED và trạng thái ON/OFF hiện tại trong database.",
        response={
            "code": 200, "message": "success",
            "result": [
                {"id": 1, "name": "LED 1", "status": "OFF", "createdAt": "2026-10-06T10:00:00"},
                {"id": 2, "name": "LED 2", "status": "ON", "createdAt": "2026-10-06T10:00:00"},
                {"id": 3, "name": "LED 3", "status": "OFF", "createdAt": "2026-10-06T10:00:00"},
            ]
        },
    )

    add_api(
        document, 9, "Điều khiển một thiết bị", "POST", "/api/devices/{id}/control",
        "Tạo lịch sử PROCESSING, gửi lệnh MQTT tới ESP8266, chờ phản hồi và cập nhật SUCCESS hoặc FAILED.",
        params=[("id", "Path", "Có", "ID LED: 1, 2 hoặc 3")],
        request_body={"action": "ON"},
        response={
            "code": 200, "message": "success",
            "result": {"id": 150, "userId": 1, "username": "admin", "deviceId": 1, "deviceName": "LED 1", "action": "ON", "status": "SUCCESS", "deviceStatus": "ON", "createdAt": "2026-10-07T15:40:00"}
        },
        errors=[
            ("Response 404 Not Found:", {"code": 404, "message": "Device not found"}),
            ("Response 504 Gateway Timeout:", {"code": 504, "message": "ESP8266 did not acknowledge the command in time"}),
        ],
        note="Cần backend, Mosquitto và ESP8266 cùng hoạt động để nhận SUCCESS.",
    )

    add_api(
        document, 10, "Bật hoặc tắt tất cả thiết bị", "POST", "/api/devices/control-all",
        "Frontend gửi một request; backend điều khiển từng LED, lưu lịch sử riêng và tổng hợp kết quả.",
        request_body={"action": "OFF"},
        response={
            "code": 200, "message": "success",
            "result": {
                "action": "OFF", "total": 3, "successCount": 3, "failedCount": 0,
                "devices": [
                    {"deviceId": 1, "deviceName": "LED 1", "action": "OFF", "status": "SUCCESS", "deviceStatus": "OFF", "message": "success"},
                    {"deviceId": 2, "deviceName": "LED 2", "action": "OFF", "status": "SUCCESS", "deviceStatus": "OFF", "message": "success"},
                    {"deviceId": 3, "deviceName": "LED 3", "action": "OFF", "status": "SUCCESS", "deviceStatus": "OFF", "message": "success"},
                ]
            }
        },
        note="Nếu một LED lỗi, API vẫn trả kết quả của các LED còn lại và failedCount lớn hơn 0.",
    )

    action_params = [
        ("deviceId", "Query", "Không", "Lọc theo ID LED"),
        ("keyword", "Query", "Không", "Tìm theo thiết bị, người thực hiện hoặc thời gian"),
        ("action", "Query", "Không", "ON hoặc OFF"),
        ("status", "Query", "Không", "SUCCESS, FAILED hoặc PROCESSING"),
        ("from", "Query", "Không", "Thời gian bắt đầu theo ISO-8601"),
        ("to", "Query", "Không", "Thời gian kết thúc theo ISO-8601"),
        ("sort", "Query", "Không", "createdAt,id kết hợp asc/desc; mặc định createdAt,desc"),
        ("page", "Query", "Không", "Trang bắt đầu từ 0"),
        ("size", "Query", "Không", "Số bản ghi mỗi trang"),
    ]
    add_api(
        document, 11, "Lấy lịch sử điều khiển", "GET", "/api/actions",
        "Tìm kiếm, lọc, sắp xếp và phân trang lịch sử bật/tắt LED tại backend.",
        params=action_params,
        response={
            "code": 200, "message": "success",
            "result": {
                "content": [{"id": 150, "userId": 1, "username": "admin", "deviceId": 1, "deviceName": "LED 1", "action": "ON", "status": "SUCCESS", "createdAt": "2026-10-07T15:40:00"}],
                "page": 0, "size": 10, "totalElements": 150, "totalPages": 15, "first": True, "last": False
            }
        },
        note="Ví dụ lọc thất bại: /api/actions?status=FAILED&page=0&size=10&sort=createdAt,desc",
    )

    add_api(
        document, 12, "Lấy chi tiết một lịch sử điều khiển", "GET", "/api/actions/{id}",
        "Lấy đúng một lần điều khiển theo ID của bản ghi action.",
        params=[("id", "Path", "Có", "ID bản ghi lịch sử điều khiển")],
        response={
            "code": 200, "message": "success",
            "result": {"id": 150, "userId": 1, "username": "admin", "deviceId": 1, "deviceName": "LED 1", "action": "ON", "status": "SUCCESS", "createdAt": "2026-10-07T15:40:00"}
        },
        errors=[("Response 404 Not Found:", {"code": 404, "message": "Action history not found"})],
    )

    profile = {
        "fullName": "Nguyễn Văn Hiển",
        "studentCode": "B23DCCN290",
        "className": "D23CNPM05",
        "email": "hiennv.b23cn290@stu.ptit.edu.vn",
        "githubUrl": "https://github.com/nguyenhien18/iot-fe.git",
        "figmaUrl": "https://www.figma.com/",
        "postmanUrl": "https://www.postman.com/",
        "reportUrl": "",
        "avatarUrl": "",
    }
    add_api(
        document, 13, "Lấy thông tin cá nhân", "GET", "/api/profile",
        "Lấy hồ sơ của tài khoản đang đăng nhập từ MySQL.",
        response={"code": 200, "message": "success", "result": profile},
        note="Backend đã có API này; trang Profile hiện tại của frontend vẫn đang hiển thị dữ liệu viết cố định trong mã nguồn.",
    )

    add_api(
        document, 14, "Cập nhật thông tin cá nhân", "PUT", "/api/profile",
        "Cập nhật toàn bộ thông tin hồ sơ của tài khoản đang đăng nhập.",
        request_body=profile,
        response={"code": 200, "message": "success", "result": profile},
        note="Dự án dùng PUT vì request gửi toàn bộ hồ sơ. Frontend Profile chưa nối API cập nhật này.",
    )

    add_heading(document, "3.16. Quy ước phản hồi và điều kiện chạy")
    rules = [
        ("Thành công", "HTTP 200, body gồm code, message và result."),
        ("Xác thực", "Trừ đăng nhập và Swagger, mọi /api/** yêu cầu JWT của tài khoản ADMIN."),
        ("Không tìm thấy", "HTTP 404 với message mô tả Sensor, Device hoặc Action không tồn tại."),
        ("Lỗi dữ liệu", "HTTP 400 khi body/query không hợp lệ."),
        ("Điều khiển thiết bị", "Cần MySQL + Backend + Mosquitto + ESP8266."),
        ("API chỉ đọc dữ liệu cũ", "Chỉ cần MySQL + Backend; không cần cắm ESP8266."),
        ("Swagger", "http://localhost:8080/swagger-ui.html"),
        ("Postman", r"D:\IOT_main\postman\IoT-Room-Monitoring.postman_collection.json"),
    ]
    add_key_value_table(document, rules)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("IoT Room Monitoring API – Nguyễn Văn Hiển")
    set_run_font(run, size=10, color=(100, 100, 100))

    document.core_properties.title = "Tài liệu REST API IoT Room Monitoring"
    document.core_properties.subject = "Mô tả 14 API của hệ thống Smart Room"
    document.core_properties.author = "Nguyễn Văn Hiển"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
